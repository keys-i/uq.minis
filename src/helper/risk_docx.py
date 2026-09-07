#!/usr/bin/env python3
"""Build an editable UQ-style event risk assessment DOCX from TOML"""

from __future__ import annotations

import re
import tomllib
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

UQ_PURPLE = "51247A"
MID_GREY = "D9D9D9"
LIGHT_GREY = "E7E7E7"
BORDER_GREY = "8A8A8A"
BLACK = "000000"
WHITE = "FFFFFF"
FONT_NAME = "Arial"


class ConfigError(ValueError):
    """Report an invalid event TOML value"""


class FormatContext(dict[str, str]):
    """Provide strict placeholder expansion for TOML text"""

    def __missing__(self, key: str) -> str:
        """Raise a useful error for an unknown placeholder"""
        raise ConfigError(f"Unknown placeholder {{{key}}} in TOML text")


@dataclass(frozen=True, slots=True)
class Schedule:
    """Hold formatted event setup, event and pack-up times"""

    setup_start: str
    setup_end: str
    event_start: str
    event_end: str
    pack_up_start: str
    pack_up_end: str

    @property
    def has_any(self) -> bool:
        """Return whether any schedule time is set"""
        return any(
            (
                self.setup_start,
                self.setup_end,
                self.event_start,
                self.event_end,
                self.pack_up_start,
                self.pack_up_end,
            )
        )

    def compact_sentence(self) -> str:
        """Summarise populated schedule ranges"""
        bits: list[str] = []
        if self.setup_start or self.setup_end:
            bits.append(f"setup {join_range(self.setup_start, self.setup_end)}")
        if self.event_start or self.event_end:
            bits.append(f"event {join_range(self.event_start, self.event_end)}")
        if self.pack_up_start or self.pack_up_end:
            bits.append(f"pack-up {join_range(self.pack_up_start, self.pack_up_end)}")
        if not bits:
            return ""
        return "Schedule: " + "; ".join(bits) + "."


@dataclass(frozen=True, slots=True)
class EventConfig:
    """Hold validated event data for risk documents"""

    source_path: Path
    output_name: str
    logo_path: Path | None
    application_title: str
    application_fields: tuple[tuple[str, str], ...]

    club_name: str
    club_short_name: str
    executives: str
    event_officer: str

    title: str
    event_date: str
    location: str
    category: str
    summary: str
    risk_scope: str
    target_audience: str
    estimated_attendance: str
    stakeholders: tuple[str, ...]
    activity: str
    entertainment: str
    additional_notes: str

    coordinator_name: str
    coordinator_role: str
    coordinator_email: str
    coordinator_phone: str

    schedule: Schedule

    indoor: bool
    seated: bool
    vehicles: bool
    temporary_structures: bool
    electrical_equipment: bool
    cash_handling: bool
    alcohol: bool
    food_provided: bool
    physical_competitions: bool
    off_site_impact: bool

    equipment_description: str
    food_description: str
    payment_method: str
    vehicle_controls: str
    temporary_structure_controls: str
    weather_controls: str
    competition_controls: str
    alcohol_controls: str

    risk_overrides: Mapping[int, Mapping[str, Any]]

    @property
    def coordinator_display(self) -> str:
        """Format coordinator details for a table cell"""
        lines = [self.coordinator_name]
        if self.coordinator_role:
            lines.append(self.coordinator_role)
        contact = " | ".join(x for x in (self.coordinator_email, self.coordinator_phone) if x)
        if contact:
            lines.append(contact)
        return "\n".join(lines)

    @property
    def description_for_cover(self) -> str:
        """Choose the cover-page event description"""
        description = self.risk_scope.strip() or self.summary.strip()
        if self.additional_notes:
            return f"{description}\n{self.additional_notes.strip()}"
        return description

    @property
    def context(self) -> FormatContext:
        """Expose event fields for TOML placeholders"""
        return FormatContext(
            club=self.club_name,
            club_short=self.club_short_name,
            executives=self.executives,
            event_officer=self.event_officer,
            event_title=self.title,
            event_date=self.event_date,
            location=self.location,
            category=self.category,
            coordinator=self.coordinator_name,
            coordinator_role=self.coordinator_role,
            coordinator_email=self.coordinator_email,
            coordinator_phone=self.coordinator_phone,
            target_audience=self.target_audience,
            attendance=self.estimated_attendance,
            activity=self.activity,
            entertainment=self.entertainment,
            equipment=self.equipment_description,
            food=self.food_description,
            payment_method=self.payment_method,
            setup_start=self.schedule.setup_start,
            setup_end=self.schedule.setup_end,
            event_start=self.schedule.event_start,
            event_end=self.schedule.event_end,
            pack_up_start=self.schedule.pack_up_start,
            pack_up_end=self.schedule.pack_up_end,
        )


@dataclass(frozen=True, slots=True)
class RiskRow:
    """Hold one completed risk assessment row"""

    number: int
    aspect: str
    hazards: tuple[str, ...]
    initial_risk: str
    controls: str
    residual_risk: str
    responsible: str


@dataclass(frozen=True, slots=True)
class RiskDefinition:
    """Describe a standard risk row and its control builder"""

    number: int
    aspect: str
    hazards: tuple[str, ...]
    builder: Callable[[EventConfig], tuple[str, str, str, str]]


# ---------------------------------------------------------------------------
# Configuration loading
# ---------------------------------------------------------------------------


def clean_text(value: Any, *, field: str = "value") -> str:
    """Render a supported TOML value as trimmed text"""
    if value is None:
        return ""
    if isinstance(value, (date, datetime, time)):
        return str(value)
    if isinstance(value, (str, int, float)):
        return str(value).strip()
    raise ConfigError(f"{field} must be a string, number, date or time")


def require_text(mapping: Mapping[str, Any], key: str, *, section: str) -> str:
    """Read a required text field from a TOML table"""
    value = clean_text(mapping.get(key), field=f"[{section}].{key}")
    if not value:
        raise ConfigError(f"Missing required [{section}].{key}")
    return value


def bool_value(mapping: Mapping[str, Any], key: str, default: bool) -> bool:
    """Read a TOML boolean with common text forms"""
    value = mapping.get(key, default)
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        normalised = value.strip().lower()
        if normalised in {"yes", "true", "1", "y"}:
            return True
        if normalised in {"no", "false", "0", "n"}:
            return False
    raise ConfigError(f"{key} must be true or false")


def format_date(value: Any) -> str:
    """Format an event date as DD/MM/YYYY"""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    raw = clean_text(value, field="event.date")
    if not raw:
        raise ConfigError("Missing required [event].date")
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).strftime("%d/%m/%Y")
        except ValueError:
            pass
    raise ConfigError("[event].date must be YYYY-MM-DD or DD/MM/YYYY")


def format_time(value: Any, *, field: str) -> str:
    """Format an optional time as HH:MM"""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    raw = clean_text(value, field=field)
    for pattern in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
        try:
            return datetime.strptime(raw, pattern).strftime("%H:%M")
        except ValueError:
            pass
    raise ConfigError(f"{field} must be a TOML time or HH:MM")


def join_range(start: str, end: str) -> str:
    """Join two optional time values into a readable range"""
    if start and end:
        return f"{start}-{end}"
    return start or end or "not specified"


def slugify(value: str) -> str:
    """Make a filesystem-safe lowercase slug"""
    value = value.lower().strip()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-") or "event"


def output_stem(value: str, *, default: str) -> str:
    """Validate a plain DOCX output filename stem"""
    stem = value.strip() or default
    if stem.lower().endswith(".docx"):
        stem = stem[:-5]
    if not stem or stem in {".", ".."} or Path(stem).name != stem or "\x00" in stem:
        raise ConfigError("[document].output_name must be a plain filename without directories")
    return stem


def display_value(value: Any, *, field: str) -> str:
    """Render a TOML value for an application summary"""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return ", ".join(clean_text(item, field=field) for item in value)
    return clean_text(value, field=field)


def read_application(data: Mapping[str, Any]) -> tuple[str, tuple[tuple[str, str], ...]]:
    """Read the optional application-summary heading and fields"""
    raw = data.get("application", {})
    if not isinstance(raw, Mapping):
        raise ConfigError("[application] must be a TOML table")
    title = clean_text(raw.get("title"), field="[application].title") or "EVENT APPLICATION SUMMARY"
    fields = raw.get("fields", {})
    if not isinstance(fields, Mapping):
        raise ConfigError("[application.fields] must be a TOML table")
    rows: list[tuple[str, str]] = []
    for label, value in fields.items():
        rendered = display_value(value, field=f"[application.fields].{label}")
        if rendered:
            rows.append((str(label).replace("_", " ").strip().title(), rendered))
    return title, tuple(rows)


def resolve_logo(config_path: Path, data: Mapping[str, Any], cli_logo: Path | None) -> Path | None:
    """Resolve a selected or configured logo file"""
    if cli_logo is not None:
        resolved = cli_logo.expanduser().resolve()
        if not resolved.is_file():
            raise ConfigError(f"Logo file does not exist: {resolved}")
        return resolved

    document = data.get("document", {})
    if not isinstance(document, Mapping):
        raise ConfigError("[document] must be a TOML table")
    logo_value = clean_text(document.get("logo"), field="[document].logo")
    if logo_value.lower() in {"none", "false", "off"}:
        return None
    if logo_value:
        candidate = Path(logo_value).expanduser()
        if not candidate.is_absolute():
            candidate = config_path.parent / candidate
        candidate = candidate.resolve()
        if not candidate.is_file():
            raise ConfigError(f"Logo file does not exist: {candidate}")
        return candidate

    return None


def build_location(event: Mapping[str, Any]) -> str:
    """Build an event location from explicit or component fields"""
    explicit = clean_text(event.get("location"), field="[event].location")
    if explicit:
        return explicit

    room = clean_text(event.get("room"), field="[event].room")
    building = clean_text(event.get("building"), field="[event].building")
    campus = clean_text(event.get("campus"), field="[event].campus")
    if not room and not building and not campus:
        raise ConfigError("Provide [event].location or at least one of room/building/campus")

    main = room
    if building:
        main = f"{main} ({building})" if main else building
    if campus:
        main = f"{main}, {campus}" if main else campus
    return main


def read_risk_overrides(data: Mapping[str, Any]) -> dict[int, Mapping[str, Any]]:
    """Read validated per-risk TOML overrides"""
    raw = data.get("risk_overrides", {})
    if raw is None:
        return {}
    if not isinstance(raw, Mapping):
        raise ConfigError("[risk_overrides] must be a TOML table")

    parsed: dict[int, Mapping[str, Any]] = {}
    for key, value in raw.items():
        try:
            number = int(key)
        except (TypeError, ValueError) as exc:
            raise ConfigError(f"Risk override key {key!r} must be a number from 1 to 22") from exc
        if not 1 <= number <= 22:
            raise ConfigError(f"Risk override {number} must be between 1 and 22")
        if not isinstance(value, Mapping):
            raise ConfigError(f"[risk_overrides.{key!r}] must be a TOML table")
        parsed[number] = value
    return parsed


def load_event_config(path: Path, cli_logo: Path | None = None) -> EventConfig:
    """Load and validate an event TOML file"""
    path = path.expanduser().resolve(strict=False)
    try:
        with path.open("rb") as handle:
            data = tomllib.load(handle)
    except OSError as exc:
        raise ConfigError(f"Cannot read {path}: {exc}") from exc
    except tomllib.TOMLDecodeError as exc:
        raise ConfigError(f"Invalid TOML in {path}: {exc}") from exc

    club = data.get("club", {})
    event = data.get("event", {})
    coordinator = data.get("coordinator", {})
    schedule_raw = data.get("schedule", {})
    catering = data.get("catering", {})
    operations = data.get("operations", {})
    document = data.get("document", {})
    application = data.get("application", {})

    for name, section in (
        ("club", club),
        ("event", event),
        ("coordinator", coordinator),
        ("schedule", schedule_raw),
        ("catering", catering),
        ("operations", operations),
        ("document", document),
        ("application", application),
    ):
        if not isinstance(section, Mapping):
            raise ConfigError(f"[{name}] must be a TOML table")

    club_name = require_text(club, "name", section="club")
    club_short_name = clean_text(club.get("short_name"), field="[club].short_name") or club_name
    executives = (
        clean_text(club.get("executives"), field="[club].executives") or f"{club_name} Executives"
    )
    event_officer = (
        clean_text(club.get("event_officer"), field="[club].event_officer")
        or f"{club_name} Events Officer"
    )

    title = require_text(event, "title", section="event")
    event_date = format_date(event.get("date"))
    location = build_location(event)
    category = clean_text(event.get("category"), field="[event].category") or "Small"
    summary = clean_text(event.get("summary"), field="[event].summary")
    if not summary:
        summary = require_text(event, "description", section="event")
    risk_scope = clean_text(event.get("risk_scope"), field="[event].risk_scope")
    target_audience = (
        clean_text(event.get("target_audience"), field="[event].target_audience") or "UQ students"
    )
    attendance = require_text(event, "estimated_attendance", section="event")

    stakeholders_raw = event.get("stakeholders", [])
    if stakeholders_raw in (None, ""):
        stakeholders: tuple[str, ...] = (
            executives,
            "event attendees",
            "UQ Union/Student Associations",
            "venue staff",
            "UQ Property and Facilities",
        )
    elif isinstance(stakeholders_raw, Sequence) and not isinstance(stakeholders_raw, (str, bytes)):
        stakeholders = tuple(
            clean_text(item, field="[event].stakeholders") for item in stakeholders_raw
        )
        stakeholders = tuple(item for item in stakeholders if item)
    else:
        raise ConfigError("[event].stakeholders must be an array of strings")

    activity = clean_text(event.get("activity"), field="[event].activity") or title
    entertainment = (
        clean_text(event.get("entertainment"), field="[event].entertainment") or activity
    )
    additional_notes = clean_text(event.get("additional_notes"), field="[event].additional_notes")

    coordinator_name = require_text(coordinator, "name", section="coordinator")
    coordinator_role = clean_text(coordinator.get("role"), field="[coordinator].role")
    coordinator_email = clean_text(coordinator.get("email"), field="[coordinator].email")
    coordinator_phone = clean_text(coordinator.get("phone"), field="[coordinator].phone")

    schedule = Schedule(
        setup_start=format_time(schedule_raw.get("setup_start"), field="[schedule].setup_start"),
        setup_end=format_time(schedule_raw.get("setup_end"), field="[schedule].setup_end"),
        event_start=format_time(schedule_raw.get("event_start"), field="[schedule].event_start"),
        event_end=format_time(schedule_raw.get("event_end"), field="[schedule].event_end"),
        pack_up_start=format_time(
            schedule_raw.get("pack_up_start"), field="[schedule].pack_up_start"
        ),
        pack_up_end=format_time(schedule_raw.get("pack_up_end"), field="[schedule].pack_up_end"),
    )

    indoor = bool_value(operations, "indoor", True)
    seated = bool_value(operations, "seated", True)
    vehicles = bool_value(operations, "vehicles", False)
    temporary_structures = bool_value(operations, "temporary_structures", False)
    electrical_equipment = bool_value(operations, "electrical_equipment", True)
    cash_handling = bool_value(operations, "cash_handling", False)
    alcohol = bool_value(operations, "alcohol", False)
    physical_competitions = bool_value(operations, "physical_competitions", False)
    off_site_impact = bool_value(operations, "off_site_impact", not indoor)

    food_provided = bool_value(catering, "provided", bool(clean_text(catering.get("description"))))
    equipment_description = (
        clean_text(operations.get("equipment"), field="[operations].equipment")
        or "existing UQ AV equipment, laptops and chargers"
    )
    food_description = clean_text(catering.get("description"), field="[catering].description")
    payment_method = (
        clean_text(operations.get("payment_method"), field="[operations].payment_method")
        or "approved cashless or pre-registration methods"
    )
    vehicle_controls = clean_text(
        operations.get("vehicle_controls"), field="[operations].vehicle_controls"
    )
    temporary_structure_controls = clean_text(
        operations.get("temporary_structure_controls"),
        field="[operations].temporary_structure_controls",
    )
    weather_controls = clean_text(
        operations.get("weather_controls"), field="[operations].weather_controls"
    )
    competition_controls = clean_text(
        operations.get("competition_controls"), field="[operations].competition_controls"
    )
    alcohol_controls = clean_text(
        operations.get("alcohol_controls"), field="[operations].alcohol_controls"
    )

    date_for_filename = datetime.strptime(event_date, "%d/%m/%Y").strftime("%Y-%m-%d")
    output_name = output_stem(
        clean_text(document.get("output_name"), field="[document].output_name"),
        default=f"{date_for_filename}-{slugify(title)}",
    )
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,119}", output_name):
        raise ConfigError(
            "[document].output_name must be 1-120 characters using only letters, "
            "numbers, dots, underscores, and hyphens"
        )
    application_title, application_fields = read_application(data)
    return EventConfig(
        source_path=path,
        output_name=output_name,
        logo_path=resolve_logo(path, data, cli_logo),
        application_title=application_title,
        application_fields=application_fields,
        club_name=club_name,
        club_short_name=club_short_name,
        executives=executives,
        event_officer=event_officer,
        title=title,
        event_date=event_date,
        location=location,
        category=category,
        summary=summary,
        risk_scope=risk_scope,
        target_audience=target_audience,
        estimated_attendance=attendance,
        stakeholders=stakeholders,
        activity=activity,
        entertainment=entertainment,
        additional_notes=additional_notes,
        coordinator_name=coordinator_name,
        coordinator_role=coordinator_role,
        coordinator_email=coordinator_email,
        coordinator_phone=coordinator_phone,
        schedule=schedule,
        indoor=indoor,
        seated=seated,
        vehicles=vehicles,
        temporary_structures=temporary_structures,
        electrical_equipment=electrical_equipment,
        cash_handling=cash_handling,
        alcohol=alcohol,
        food_provided=food_provided,
        physical_competitions=physical_competitions,
        off_site_impact=off_site_impact,
        equipment_description=equipment_description,
        food_description=food_description,
        payment_method=payment_method,
        vehicle_controls=vehicle_controls,
        temporary_structure_controls=temporary_structure_controls,
        weather_controls=weather_controls,
        competition_controls=competition_controls,
        alcohol_controls=alcohol_controls,
        risk_overrides=read_risk_overrides(data),
    )


# ---------------------------------------------------------------------------
# Risk rows
# ---------------------------------------------------------------------------


def access_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for access and egress"""
    if cfg.vehicles:
        vehicle_text = (
            cfg.vehicle_controls
            or "Any vehicle movement or delivery activity will be supervised and separated from pedestrians."
        )
        initial = "Medium"
    else:
        vehicle_text = "No vehicles are expected to be involved in the event."
        initial = "Low"
    controls = (
        f"The event will be held at {cfg.location} using normal pedestrian access and marked exits. "
        "Walkways, doors and emergency exits will be kept clear at all times. " + vehicle_text
    )
    return initial, controls, "Low", cfg.executives


def setup_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for setup and pack-up"""
    timing = ""
    if cfg.schedule.setup_start or cfg.schedule.pack_up_end:
        timing = (
            f" Setup is scheduled for {join_range(cfg.schedule.setup_start, cfg.schedule.setup_end)} and "
            f"pack-up for {join_range(cfg.schedule.pack_up_start, cfg.schedule.pack_up_end)}."
        )
    controls = (
        "Only light setup is required and the existing room layout will be retained where possible. "
        "Furniture and materials will be moved carefully, and bags, boxes and cables will be kept out of walkways."
        + timing
    )
    return "Low-Medium", controls, "Low", cfg.executives


def structures_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for temporary structures"""
    if not cfg.temporary_structures:
        return (
            "N/A",
            "No temporary structures, marquees, stages, barriers or large temporary equipment will be used.",
            "N/A",
            "N/A",
        )
    controls = cfg.temporary_structure_controls or (
        "Temporary structures will be installed by competent persons, checked before use and positioned clear of exits and access ways."
    )
    return "Medium", controls, "Low", cfg.executives


def weather_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for weather exposure"""
    if cfg.indoor:
        controls = "The event is indoors, so weather exposure is minimal. Attendees will use normal campus paths and buildings when arriving and leaving."
        return "Low", controls, "Low", cfg.event_officer
    controls = cfg.weather_controls or (
        "Weather conditions will be checked before the event. Shelter, water and a cancellation or relocation plan will be available if conditions become unsafe."
    )
    return "Medium", controls, "Low", cfg.event_officer


def electrical_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for electrical equipment"""
    if not cfg.electrical_equipment:
        return "N/A", "No electrically powered equipment is planned for the event.", "N/A", "N/A"
    equipment = cfg.equipment_description
    equipment = equipment[:1].upper() + equipment[1:] if equipment else "Electrical equipment"
    controls = (
        f"{equipment} will be used where required. Equipment and leads will be checked for visible damage. "
        "Cables will be routed away from walkways and kept clear of food and drinks."
    )
    return "Low-Medium", controls, "Low", cfg.event_officer


def traffic_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for vehicle movement"""
    if not cfg.vehicles:
        return (
            "N/A",
            "No vehicle access, road closures, deliveries or traffic control are required.",
            "N/A",
            "N/A",
        )
    controls = cfg.vehicle_controls or (
        "Vehicle access and deliveries will use approved campus routes, be scheduled outside attendee arrival periods and be supervised by event staff."
    )
    return "Medium", controls, "Low", cfg.event_officer


def crowd_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for attendance and crowd movement"""
    posture = "seated and structured" if cfg.seated else "managed by event staff"
    controls = (
        f"Attendance is expected to be {cfg.estimated_attendance} people and will not exceed the booked room capacity. "
        f"The event will be {posture}. Exits and aisles will remain clear, and attendance will be monitored."
    )
    return "Low-Medium", controls, "Low", cfg.executives


def waste_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for event waste"""
    catering = ""
    if cfg.food_provided and cfg.food_description:
        catering = f" Waste from {cfg.food_description} will be placed in available bins."
    controls = (
        "Waste is expected to be limited. Attendees will use available bins, and event staff will inspect the room after the event and remove rubbish or leftover materials."
        + catering
    )
    return "Low", controls, "Low", cfg.executives


def amenities_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for venue amenities"""
    controls = "The event is held at a UQ venue with existing toilet facilities nearby. The expected attendance is small enough that amenities are unlikely to be overloaded."
    return "Low", controls, "Low", cfg.event_officer


def wellness_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for attendee and staff wellbeing"""
    setting = (
        "seated, indoors and low-risk" if cfg.indoor and cfg.seated else "low-risk and supervised"
    )
    controls = f"The event is {setting}. Attendees may take breaks or leave if needed. {cfg.executives} will monitor for illness, distress or discomfort and contact UQ Security or emergency services if required."
    return "Low-Medium", controls, "Low", cfg.executives


def fire_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build fire safety controls"""
    controls = (
        "The event will use an existing UQ venue with fire alarms, marked exits and evacuation procedures. "
        "No open flames, candles or unapproved heat-producing equipment will be used, and exits will remain clear."
    )
    return "Low-Medium", controls, "Low", cfg.executives


def entertainment_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for entertainment activities"""
    controls = (
        f"Entertainment and activities are limited to {cfg.entertainment}. Sound will remain at a reasonable level. "
        "No high-noise or disruptive entertainment is planned."
    )
    return "Low", controls, "Low", cfg.event_officer


def cash_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for cash handling"""
    if not cfg.cash_handling:
        return (
            "N/A",
            f"No cash handling is expected. Any payment, if required, will use {cfg.payment_method}.",
            "N/A",
            "N/A",
        )
    controls = (
        f"Cash will be handled only by authorised {cfg.club_short_name} officers, counted by two people and secured promptly. "
        f"Cashless payment through {cfg.payment_method} will be preferred."
    )
    return "Medium", controls, "Low", cfg.executives


def security_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build security controls"""
    controls = (
        f"{cfg.executives} will be present throughout the event and attendance will be monitored. "
        "Disruptive, unsafe or unauthorised behaviour will be addressed promptly and escalated to UQ Security if needed."
    )
    return "Low", controls, "Low", cfg.executives


def alcohol_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for alcohol service"""
    if not cfg.alcohol:
        return "N/A", "No alcohol will be served or permitted at the event.", "N/A", "N/A"
    controls = cfg.alcohol_controls or (
        "Alcohol service will proceed only with all required UQ and liquor approvals, responsible service controls, age verification, water and food availability, and a plan for intoxicated or unwell patrons."
    )
    return "High", controls, "Medium", cfg.executives


def food_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for food service"""
    if not cfg.food_provided:
        return "N/A", "No food will be prepared, served or distributed at the event.", "N/A", "N/A"
    food = cfg.food_description or "light refreshments"
    controls = (
        f"Food service is limited to {food}. Food will be commercially prepared or handled safely, allergens will be identified where possible, "
        "food will be kept clear of electrical equipment, and spills will be cleaned promptly."
    )
    return "Low-Medium", controls, "Low", cfg.executives


def competition_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for competitions"""
    if not cfg.physical_competitions:
        controls = f"Planned activity is limited to {cfg.activity}. No physical, hazardous or impromptu competitions will be permitted."
        return "Low", controls, "Low", cfg.event_officer
    controls = cfg.competition_controls or (
        "Competition rules, participant limits and a clear activity area will be established. Event staff will supervise and stop any unsafe activity."
    )
    return "Medium", controls, "Low", cfg.event_officer


def incident_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build incident-management controls"""
    controls = (
        f"{cfg.executives} will be present and incidents will be reported to {cfg.coordinator_name}. "
        "UQ emergency procedures will be followed, with UQ Security or emergency services contacted when required."
    )
    return "Low-Medium", controls, "Low", cfg.executives


def consequential_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for serious conduct risks"""
    controls = "Respectful behaviour expectations apply throughout the event. Harassment, conflict, unwanted conduct or any safety concern will be taken seriously and escalated to UQ Security, UQ Union or relevant support services where required."
    return "Low-Medium", controls, "Low", cfg.executives


def assets_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for campus assets"""
    controls = "Existing room furniture, fixtures and equipment will be used appropriately. The venue will be checked after the event and returned to its original condition."
    return "Low", controls, "Low", cfg.executives


def offsite_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for off-site effects"""
    if not cfg.off_site_impact:
        return (
            "N/A",
            "The event is small-scale and contained within the booked venue, so it is unlikely to affect neighbours or off-site areas.",
            "N/A",
            "N/A",
        )
    controls = "Noise, attendee movement and access will be monitored to minimise impacts beyond the event area. Complaints will be addressed promptly."
    return "Low-Medium", controls, "Low", cfg.event_officer


def reputation_controls(cfg: EventConfig) -> tuple[str, str, str, str]:
    """Build controls for reputational risks"""
    controls = f"The event will be run respectfully and in line with UQ and {cfg.club_name} expectations. Complaints, misconduct or negative feedback will be addressed promptly and professionally."
    return "Low", controls, "Low", cfg.executives


RISK_DEFINITIONS: tuple[RiskDefinition, ...] = (
    RiskDefinition(
        1,
        "Access / Egress to and from site",
        (
            "Vehicle movements - potential for persons to be injured by reversing vehicles",
            "Attendee movement - ability to exit safely and quickly",
            "Other:",
        ),
        access_controls,
    ),
    RiskDefinition(
        2,
        "Set Up and Pack Up",
        (
            "Slips, trips and falls",
            "Pinch points / caught between crush points",
            "Strains and sprains from manual handling exertion",
            "Incorrect set up of equipment",
            "Other:",
        ),
        setup_controls,
    ),
    RiskDefinition(
        3,
        "Temporary Structures",
        (
            "Collapse",
            "Trip hazards",
            "Obstruction of walkways, exits or emergency access",
            "Other:",
        ),
        structures_controls,
    ),
    RiskDefinition(
        4,
        "Adverse Weather",
        (
            "Harm to patrons from cold or wet conditions if no shelter",
            "High winds",
            "Other:",
        ),
        weather_controls,
    ),
    RiskDefinition(
        5,
        "Electrically Powered Equipment",
        (
            "Electric shock",
            "Trip hazards from electric cords",
            "Incorrectly placed equipment",
            "Other:",
        ),
        electrical_controls,
    ),
    RiskDefinition(
        6,
        "Traffic Control",
        ("Vehicles on site before, during or after the event",),
        traffic_controls,
    ),
    RiskDefinition(
        7,
        "Number of Patrons (Crowd Control)",
        (
            "Panic",
            "Protestors",
            "Bomb threat",
            "Trampling (crowd related)",
            "Crush (crowd related movement)",
            "Excessive number of people at venue",
            "Other:",
        ),
        crowd_controls,
    ),
    RiskDefinition(
        8, "Waste Management", ("Overflowing waste containers", "Other:"), waste_controls
    ),
    RiskDefinition(
        9,
        "Provision of Amenities",
        (
            "Public urination (indecent exposure)",
            "Excessive queues for toilets",
            "Contamination of toilet amenities",
            "Overload of toilet amenities",
            "Other:",
        ),
        amenities_controls,
    ),
    RiskDefinition(
        10,
        "Patron and Staff (worker) wellness",
        ("Heat exhaustion", "Dehydration", "Stress, anxiety or other psychosocial risk", "Other:"),
        wellness_controls,
    ),
    RiskDefinition(
        11,
        "Fire Safety",
        ("Fire on site", "Patron and staff safety during a fire event", "Other:"),
        fire_controls,
    ),
    RiskDefinition(
        12,
        "Provision of Entertainment",
        ("High noise levels leading to public complaint", "Other:"),
        entertainment_controls,
    ),
    RiskDefinition(13, "Cash Handling", ("Theft", "Hygiene", "Other:"), cash_controls),
    RiskDefinition(
        14,
        "Security Matters",
        (
            "Forced entry / breach of perimeter",
            "Protestors",
            "Excessive number of patrons at venue",
            "Undesired patron entry",
            "Protection of attendees",
            "Drug use on site",
            "Other:",
        ),
        security_controls,
    ),
    RiskDefinition(
        15,
        "Service of Alcohol",
        (
            "Breach of liquor licence",
            "Excessive intoxication",
            "Sick patrons",
            "Persons requiring medical attention",
            "Supply to minors",
            "Cross contamination during serving of premixed drinks",
            "Other:",
        ),
        alcohol_controls,
    ),
    RiskDefinition(
        16,
        "Service of Food",
        ("Food poisoning through inadequate food preparation, handling and control", "Other:"),
        food_controls,
    ),
    RiskDefinition(
        17,
        "Competitions",
        ("Injury during planned or impromptu competitions", "Other:"),
        competition_controls,
    ),
    RiskDefinition(
        18,
        "On Site Incident Management",
        (
            "Communication of incident",
            "Emergency alarms and protocols",
            "Emergency Control Organisation",
            "Other:",
        ),
        incident_controls,
    ),
    RiskDefinition(
        19,
        "Consequential Outcomes",
        ("Sexual assault", "Conflict or fights", "Other:"),
        consequential_controls,
    ),
    RiskDefinition(
        20,
        "Campus Assets",
        ("Damage to flora and fauna", "Damage to buildings and structures", "Other:"),
        assets_controls,
    ),
    RiskDefinition(
        21,
        "Off Site Matters",
        (
            "Public complaints",
            "Neighbours impacted by noise levels",
            "Neighbours impacted by intruders",
            "Other:",
        ),
        offsite_controls,
    ),
    RiskDefinition(
        22,
        "Reputational risks",
        ("Loss of reputation", "Misconduct", "Negative feedback", "Other:"),
        reputation_controls,
    ),
)


def expand_text(value: Any, cfg: EventConfig, *, field: str) -> str:
    """Expand event placeholders in an override value"""
    text = clean_text(value, field=field)
    return text.format_map(cfg.context) if text else ""


def build_risks(cfg: EventConfig) -> list[RiskRow]:
    """Build standard risk rows with TOML overrides"""
    rows: list[RiskRow] = []
    for definition in RISK_DEFINITIONS:
        initial, controls, residual, responsible = definition.builder(cfg)
        aspect = definition.aspect
        hazards = definition.hazards
        override = cfg.risk_overrides.get(definition.number, {})

        if "aspect" in override:
            aspect = expand_text(override["aspect"], cfg, field=f"risk {definition.number} aspect")
        if "hazards" in override:
            raw_hazards = override["hazards"]
            if not isinstance(raw_hazards, Sequence) or isinstance(raw_hazards, (str, bytes)):
                raise ConfigError(f"Risk {definition.number} hazards must be an array of strings")
            hazards = tuple(
                expand_text(item, cfg, field=f"risk {definition.number} hazards")
                for item in raw_hazards
            )
        if "initial_risk" in override:
            initial = expand_text(
                override["initial_risk"], cfg, field=f"risk {definition.number} initial_risk"
            )
        if "controls" in override:
            controls = expand_text(
                override["controls"], cfg, field=f"risk {definition.number} controls"
            )
        if "residual_risk" in override:
            residual = expand_text(
                override["residual_risk"], cfg, field=f"risk {definition.number} residual_risk"
            )
        if "responsible" in override:
            responsible = expand_text(
                override["responsible"], cfg, field=f"risk {definition.number} responsible"
            )

        rows.append(
            RiskRow(
                number=definition.number,
                aspect=aspect,
                hazards=hazards,
                initial_risk=initial,
                controls=controls,
                residual_risk=residual,
                responsible=responsible,
            )
        )
    return rows


# ---------------------------------------------------------------------------
# DOCX helpers
# ---------------------------------------------------------------------------


def set_run_font(
    run: Any,
    *,
    size: float,
    bold: bool | None = None,
    italic: bool | None = None,
    colour: str = BLACK,
    name: str = FONT_NAME,
) -> None:
    """Apply the document font settings to a run"""
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(colour)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell: Any, fill: str) -> None:
    """Fill a table cell with a Word colour"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(
    cell: Any, *, top: int = 45, start: int = 55, bottom: int = 45, end: int = 55
) -> None:
    """Set Word table-cell margins in twips"""
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = f"w:{edge}"
        element = tc_mar.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell: Any, width_inches: float) -> None:
    """Set a table-cell width in inches"""
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_fixed_widths(table: Any, widths_inches: Sequence[float]) -> None:
    """Force a Word table to use the supplied grid widths"""
    if len(table.columns) != len(widths_inches):
        raise ValueError("Column-width count does not match table column count")

    table.autofit = False
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    total_twips = sum(int(Inches(width).twips) for width in widths_inches)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")

    grid_columns = list(table._tbl.tblGrid.gridCol_lst)
    if len(grid_columns) != len(widths_inches):
        raise ValueError("Unexpected DOCX table grid")
    for grid_col, width in zip(grid_columns, widths_inches):
        grid_col.set(qn("w:w"), str(int(Inches(width).twips)))

    for column, width in zip(table.columns, widths_inches):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_inches):
            set_cell_width(cell, width)


def set_table_borders(table: Any, *, colour: str = BORDER_GREY, size: int = 4) -> None:
    """Apply single borders to every table edge"""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), colour)


def remove_table_borders(table: Any) -> None:
    """Hide every table border"""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)


def prevent_row_split(row: Any) -> None:
    """Keep a Word table row together across pages"""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_repeat_table_header(row: Any) -> None:
    """Repeat a table row at each page break"""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_paragraph_layout(
    paragraph: Any,
    *,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
    before: float = 0,
    after: float = 0,
    line: float | None = None,
    keep_with_next: bool = False,
) -> None:
    """Apply paragraph alignment, spacing and line settings"""
    if alignment is not None:
        paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.keep_with_next = keep_with_next
    if line is not None:
        fmt.line_spacing = Pt(line)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def clear_cell(cell: Any) -> None:
    """Empty a table cell and initialise its paragraph"""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_layout(paragraph, line=8.5)


def add_text_to_cell(
    cell: Any,
    text: str,
    *,
    size: float = 8.0,
    bold: bool = False,
    italic: bool = False,
    colour: str = BLACK,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    line: float | None = None,
) -> None:
    """Replace a table cell with formatted multi-line text"""
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    set_paragraph_layout(paragraph, alignment=alignment, line=line or max(size + 0.8, 8.0))
    lines = text.split("\n") if text else [""]
    for idx, item in enumerate(lines):
        if idx:
            paragraph.add_run().add_break()
        run = paragraph.add_run(item)
        set_run_font(run, size=size, bold=bold, italic=italic, colour=colour)


def add_page_number(paragraph: Any) -> None:
    """Insert a dynamic PAGE field into a paragraph"""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_layout(paragraph, line=9.0)
    run = paragraph.add_run()
    set_run_font(run, size=8.0)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    for element in (fld_char_begin, instr_text, fld_char_separate, display, fld_char_end):
        run._r.append(element)


def apply_document_defaults(doc: DocumentType) -> None:
    """Apply page, font and footer defaults to a document"""
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)
    section.header_distance = Inches(0.20)
    section.footer_distance = Inches(0.20)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(8.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""
    footer = section.footer
    footer.paragraphs[0].text = ""
    add_page_number(footer.paragraphs[0])


def set_core_properties(doc: DocumentType, cfg: EventConfig) -> None:
    """Set searchable document metadata from event data"""
    props = doc.core_properties
    props.title = f"Event Management Risk Assessment - {cfg.title}"
    props.subject = cfg.summary
    props.author = cfg.coordinator_name
    props.keywords = f"risk assessment, event, {cfg.club_name}"
    props.comments = f"Generated from {cfg.source_path.name}"


# ---------------------------------------------------------------------------
# DOCX construction
# ---------------------------------------------------------------------------


def add_cover_page(doc: DocumentType, cfg: EventConfig) -> None:
    """Add the editable UQ risk-assessment cover page"""
    usable_width = 10.69
    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_widths(header, (4.2, usable_width - 4.2))
    remove_table_borders(header)
    header.cell(0, 0).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    header.cell(0, 1).vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(header.cell(0, 0), top=0, start=0, bottom=0, end=0)
    set_cell_margins(header.cell(0, 1), top=0, start=0, bottom=0, end=0)

    header.rows[0].height = Inches(0.96)
    header.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    left = header.cell(0, 0).paragraphs[0]
    set_paragraph_layout(left, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if cfg.logo_path and cfg.logo_path.is_file():
        left.add_run().add_picture(str(cfg.logo_path), width=Inches(2.30))
    else:
        run = left.add_run("THE UNIVERSITY OF QUEENSLAND\nAUSTRALIA")
        set_run_font(run, size=12.0, bold=True, colour=UQ_PURPLE)

    right = header.cell(0, 1).paragraphs[0]
    set_paragraph_layout(right, alignment=WD_ALIGN_PARAGRAPH.RIGHT, line=10.0)
    run = right.add_run("Property and Facilities Division")
    set_run_font(run, size=8.5, bold=True, colour=UQ_PURPLE)
    run.add_break()
    run2 = right.add_run("Form")
    set_run_font(run2, size=8.5, bold=True, colour=UQ_PURPLE)

    spacer = doc.add_paragraph()
    set_paragraph_layout(spacer, after=1, line=2.0)

    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_widths(title_table, (usable_width,))
    remove_table_borders(title_table)
    shade_cell(title_table.cell(0, 0), UQ_PURPLE)
    set_cell_margins(title_table.cell(0, 0), top=35, start=55, bottom=35, end=55)
    add_text_to_cell(
        title_table.cell(0, 0),
        "EVENT MANAGEMENT RISK ASSESSMENT",
        size=10.0,
        bold=True,
        colour=WHITE,
        line=10.5,
    )

    p = doc.add_paragraph()
    set_paragraph_layout(p, before=4, after=3, line=10.5)
    run = p.add_run("Summary of Event Management Risk Assessment. This cover section is ")
    set_run_font(run, size=8.5)
    run = p.add_run("mandatory")
    set_run_font(run, size=8.5, bold=True)
    run = p.add_run(" for all Risk Management plans submitted:")
    set_run_font(run, size=8.5)

    # Four-column grid, with merges used for two-column rows.
    widths = (2.35, 4.95, 2.15, 1.24)
    table = doc.add_table(rows=10, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_widths(table, widths)
    set_table_borders(table)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_margins(cell, top=45, start=65, bottom=45, end=65)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    normal_rows = [
        (0, "Event Title:", cfg.title),
        (1, "Event Date:", cfg.event_date),
        (2, "Proposed Location:", cfg.location),
        (3, "Event Category:", cfg.category),
        (4, "Event Coordinator Name:", cfg.coordinator_display),
        (5, "Organisation:", cfg.club_name),
    ]
    for row_idx, label, value in normal_rows:
        value_cell = table.cell(row_idx, 1).merge(table.cell(row_idx, 3))
        shade_cell(table.cell(row_idx, 0), MID_GREY)
        add_text_to_cell(table.cell(row_idx, 0), label, size=8.3, bold=True, line=9.0)
        add_text_to_cell(value_cell, value, size=8.3, line=9.0)

    description_cell = table.cell(6, 1).merge(table.cell(6, 3))
    shade_cell(table.cell(6, 0), MID_GREY)
    clear_cell(table.cell(6, 0))
    label_p = table.cell(6, 0).paragraphs[0]
    set_paragraph_layout(label_p, line=9.0)
    run = label_p.add_run("Description of the Event:")
    set_run_font(run, size=8.2, bold=True)
    run.add_break()
    run = label_p.add_run("(Provide context and scope for the risk assessment)")
    set_run_font(run, size=7.6, italic=True)
    add_text_to_cell(description_cell, cfg.description_for_cover, size=7.75, line=8.4)
    description_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    table.cell(6, 0).vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    shade_cell(table.cell(7, 0), MID_GREY)
    shade_cell(table.cell(7, 2), MID_GREY)
    add_text_to_cell(table.cell(7, 0), "Target audience:", size=8.2, bold=True, line=9.0)
    add_text_to_cell(table.cell(7, 1), cfg.target_audience, size=8.2, line=9.0)
    add_text_to_cell(table.cell(7, 2), "Estimated total attendance:", size=8.0, bold=True, line=8.8)
    add_text_to_cell(
        table.cell(7, 3),
        cfg.estimated_attendance,
        size=8.2,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line=9.0,
    )

    stakeholders_cell = table.cell(8, 1).merge(table.cell(8, 3))
    shade_cell(table.cell(8, 0), MID_GREY)
    add_text_to_cell(table.cell(8, 0), "Stakeholders:", size=8.2, bold=True, line=9.0)
    add_text_to_cell(stakeholders_cell, ", ".join(cfg.stakeholders), size=7.8, line=8.5)

    # Organisation review row intentionally blank.
    shade_cell(table.cell(9, 0), MID_GREY)
    shade_cell(table.cell(9, 2), MID_GREY)
    add_text_to_cell(
        table.cell(9, 0),
        "Reviewed by Organisation\nSafety/OHS Coordinator:",
        size=7.8,
        bold=True,
        line=8.3,
    )
    add_text_to_cell(table.cell(9, 1), "", size=8.0)
    add_text_to_cell(table.cell(9, 2), "Date:", size=8.0, bold=True, line=8.5)
    add_text_to_cell(table.cell(9, 3), "Approved:\nComments:", size=7.8, bold=True, line=8.3)

    # Keep page 1 visually balanced and prevent a stray trailing paragraph from
    # acquiring visible spacing.
    tail = doc.add_paragraph()
    set_paragraph_layout(tail, line=1.0)


def add_risk_table_header(table: Any) -> None:
    """Add the two-row heading for a risk table"""
    top = table.rows[0]
    left = table.cell(0, 0).merge(table.cell(0, 2))
    right = table.cell(0, 3).merge(table.cell(0, 6))
    for cell in (left, right):
        shade_cell(cell, MID_GREY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    add_text_to_cell(left, "RISK IDENTIFICATION", size=8.3, bold=True, line=9.0)
    add_text_to_cell(right, "DECISIONS and CONTROLS - ACTION PLAN", size=8.3, bold=True, line=9.0)
    prevent_row_split(top)
    set_repeat_table_header(top)

    headings = (
        "#",
        "ASPECT",
        "LOSS EXPOSURE, or POTENTIAL\nADVERSE OUTCOME, or HAZARD, etc.",
        "Initial Risk\n(Before Controls)",
        "RECOMMENDED CONTROLS, or\nEXISTING CONTROLS\n(How will the risk be managed)",
        "Planned Residual Risk\n(After Controls)",
        "Responsible Person(s)",
    )
    row = table.rows[1]
    for idx, heading in enumerate(headings):
        cell = row.cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        clear_cell(cell)
        p = cell.paragraphs[0]
        align = WD_ALIGN_PARAGRAPH.CENTER if idx in {0, 3, 5, 6} else WD_ALIGN_PARAGRAPH.LEFT
        set_paragraph_layout(p, alignment=align, line=7.9)
        lines = heading.split("\n")
        for line_idx, line in enumerate(lines):
            if line_idx:
                p.add_run().add_break()
            italic = line.startswith("(")
            run = p.add_run(line)
            set_run_font(run, size=7.2 if italic else 7.5, bold=not italic, italic=italic)
    prevent_row_split(row)
    set_repeat_table_header(row)


def estimated_row_height(row: RiskRow) -> float:
    """Estimate a readable minimum risk-row height in inches"""
    # Conservative minimum height in inches. Word may grow the row if required.
    hazard_lines = sum(max(1, (len(line) + 34) // 35) for line in row.hazards)
    control_lines = max(2, (len(row.controls) + 47) // 48)
    aspect_lines = max(1, (len(row.aspect) + 18) // 19)
    lines = max(hazard_lines, control_lines, aspect_lines)
    return max(0.55, min(1.55, 0.13 * lines + 0.18))


def add_risk_page(doc: DocumentType, rows: Sequence[RiskRow]) -> None:
    """Add a paginated risk-assessment table"""
    widths = (0.32, 1.47, 2.73, 0.72, 2.76, 0.78, 1.91)
    table = doc.add_table(rows=2, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_widths(table, widths)
    set_table_borders(table)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_margins(cell, top=35, start=45, bottom=35, end=45)
    add_risk_table_header(table)

    for risk in rows:
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell, top=35, start=45, bottom=35, end=45)
        row = table.rows[-1]
        prevent_row_split(row)
        row.height = Inches(estimated_row_height(risk))
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

        add_text_to_cell(
            cells[0], str(risk.number), size=7.2, alignment=WD_ALIGN_PARAGRAPH.CENTER, line=7.8
        )
        add_text_to_cell(cells[1], risk.aspect, size=7.25, bold=True, line=7.9)
        add_text_to_cell(cells[2], "\n".join(risk.hazards), size=7.0, line=7.65)
        add_text_to_cell(
            cells[3], risk.initial_risk, size=7.2, alignment=WD_ALIGN_PARAGRAPH.CENTER, line=7.8
        )
        add_text_to_cell(cells[4], risk.controls, size=7.0, line=7.65)
        add_text_to_cell(
            cells[5], risk.residual_risk, size=7.2, alignment=WD_ALIGN_PARAGRAPH.CENTER, line=7.8
        )
        add_text_to_cell(
            cells[6], risk.responsible, size=7.0, alignment=WD_ALIGN_PARAGRAPH.CENTER, line=7.65
        )

        cells[0].vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cells[1].vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cells[2].vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cells[3].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cells[4].vertical_alignment = WD_ALIGN_VERTICAL.TOP
        cells[5].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cells[6].vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_uq_only_page(doc: DocumentType) -> None:
    """Add the UQ review and approval page"""
    usable_width = 10.69
    table = doc.add_table(rows=4, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (2.35, 2.70, 0.90, 1.35, 1.65, 1.74)
    set_table_fixed_widths(table, widths)
    set_table_borders(table)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell, top=55, start=65, bottom=55, end=65)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    header = table.cell(0, 0).merge(table.cell(0, 5))
    shade_cell(header, BLACK)
    add_text_to_cell(
        header,
        "THIS SECTION TO BE COMPLETED BY THE UNIVERSITY OF QUEENSLAND ONLY",
        size=8.6,
        bold=True,
        colour=WHITE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        line=9.2,
    )

    shade_cell(table.cell(1, 0), MID_GREY)
    shade_cell(table.cell(1, 2), MID_GREY)
    shade_cell(table.cell(1, 4), MID_GREY)
    add_text_to_cell(table.cell(1, 0), "Received by Security:", size=8.0, bold=True)
    add_text_to_cell(table.cell(1, 1), "", size=8.0)
    add_text_to_cell(table.cell(1, 2), "Date:", size=8.0, bold=True)
    add_text_to_cell(table.cell(1, 3), "", size=8.0)
    sufficient = table.cell(1, 4).merge(table.cell(1, 5))
    add_text_to_cell(sufficient, "Sufficient:   Y   N", size=8.0, bold=True)

    shade_cell(table.cell(2, 0), MID_GREY)
    add_text_to_cell(table.cell(2, 0), "Event Risk Category:", size=8.0, bold=True)
    add_text_to_cell(table.cell(2, 1), "Minor", size=8.0, bold=True)
    add_text_to_cell(table.cell(2, 2), "Major", size=8.0, bold=True)
    blank = table.cell(2, 3).merge(table.cell(2, 5))
    add_text_to_cell(blank, "", size=8.0)

    shade_cell(table.cell(3, 0), MID_GREY)
    shade_cell(table.cell(3, 2), MID_GREY)
    shade_cell(table.cell(3, 4), MID_GREY)
    add_text_to_cell(table.cell(3, 0), "Received by OHS:", size=8.0, bold=True)
    add_text_to_cell(table.cell(3, 1), "", size=8.0)
    add_text_to_cell(table.cell(3, 2), "Date:", size=8.0, bold=True)
    add_text_to_cell(table.cell(3, 3), "", size=8.0)
    sufficient = table.cell(3, 4).merge(table.cell(3, 5))
    add_text_to_cell(sufficient, "Sufficient:   Y   N", size=8.0, bold=True)

    if abs(sum(widths) - usable_width) >= 0.1:
        raise ValueError("Approval table widths exceed the page width")


def add_form_content(doc: DocumentType, cfg: EventConfig) -> None:
    """Add the full risk-assessment form content"""
    add_cover_page(doc, cfg)
    risks = build_risks(cfg)
    for group in (risks[0:4], risks[4:9], risks[9:14], risks[14:18], risks[18:22]):
        doc.add_page_break()
        add_risk_page(doc, group)
    doc.add_page_break()
    add_uq_only_page(doc)


def yes_no(value: bool) -> str:
    """Render a boolean as Yes or No"""
    return "Yes" if value else "No"


def add_application_summary_page(doc: DocumentType, cfg: EventConfig) -> None:
    """Add the optional event application summary"""
    usable_width = 10.69
    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_widths(header, (4.2, usable_width - 4.2))
    remove_table_borders(header)
    for cell in header.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left = header.cell(0, 0).paragraphs[0]
    set_paragraph_layout(left, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if cfg.logo_path and cfg.logo_path.is_file():
        left.add_run().add_picture(str(cfg.logo_path), width=Inches(2.05))
    else:
        run = left.add_run("THE UNIVERSITY OF QUEENSLAND\nAUSTRALIA")
        set_run_font(run, size=11.0, bold=True, colour=UQ_PURPLE)

    right = header.cell(0, 1).paragraphs[0]
    set_paragraph_layout(right, alignment=WD_ALIGN_PARAGRAPH.RIGHT, line=9.4)
    run = right.add_run(cfg.club_name)
    set_run_font(run, size=9.0, bold=True, colour=UQ_PURPLE)
    run.add_break()
    run = right.add_run("Event document pack")
    set_run_font(run, size=8.0, colour=UQ_PURPLE)

    title = doc.add_table(rows=1, cols=1)
    title.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_widths(title, (usable_width,))
    remove_table_borders(title)
    shade_cell(title.cell(0, 0), UQ_PURPLE)
    set_cell_margins(title.cell(0, 0), top=38, start=55, bottom=38, end=55)
    add_text_to_cell(
        title.cell(0, 0),
        cfg.application_title.upper(),
        size=10.0,
        bold=True,
        colour=WHITE,
        line=10.5,
    )

    summary = doc.add_paragraph()
    set_paragraph_layout(summary, before=4, after=4, line=9.0)
    run = summary.add_run(cfg.summary)
    set_run_font(run, size=8.0, bold=True)

    extra = list(cfg.application_fields)
    table = doc.add_table(rows=12 + bool(cfg.additional_notes) + len(extra), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (1.62, 3.73, 1.62, 3.72)
    set_table_fixed_widths(table, widths)
    set_table_borders(table)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell, top=30, start=48, bottom=30, end=48)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    def pair(
        row: int, left_label: str, left_value: str, right_label: str, right_value: str
    ) -> None:
        """Fill one two-column summary row"""
        for cell_index in (0, 2):
            shade_cell(table.cell(row, cell_index), MID_GREY)
        add_text_to_cell(table.cell(row, 0), left_label, size=7.25, bold=True, line=7.8)
        add_text_to_cell(table.cell(row, 1), left_value, size=7.25, line=7.8)
        add_text_to_cell(table.cell(row, 2), right_label, size=7.25, bold=True, line=7.8)
        add_text_to_cell(table.cell(row, 3), right_value, size=7.25, line=7.8)

    def full(row: int, label: str, value: str) -> None:
        """Fill one full-width summary row"""
        value_cell = table.cell(row, 1).merge(table.cell(row, 3))
        shade_cell(table.cell(row, 0), MID_GREY)
        add_text_to_cell(table.cell(row, 0), label, size=7.2, bold=True, line=7.7)
        add_text_to_cell(value_cell, value, size=7.05, line=7.65)
        value_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    pair(0, "Event title", cfg.title, "Date", cfg.event_date)
    pair(1, "Organisation", cfg.club_name, "Category", cfg.category)
    pair(2, "Location", cfg.location, "Attendance", cfg.estimated_attendance)
    pair(3, "Target audience", cfg.target_audience, "Coordinator", cfg.coordinator_name)
    pair(
        4,
        "Coordinator role",
        cfg.coordinator_role or "Not specified",
        "Phone",
        cfg.coordinator_phone or "Not specified",
    )
    full(5, "Email", cfg.coordinator_email or "Not specified")
    full(
        6, "Schedule", cfg.schedule.compact_sentence().removeprefix("Schedule: ") or "Not specified"
    )
    full(7, "Risk scope", cfg.risk_scope or cfg.summary)
    full(8, "Activity", cfg.activity)
    full(9, "Entertainment", cfg.entertainment)
    catering = cfg.food_description if cfg.food_provided else "No food service planned"
    full(
        10,
        "Catering",
        f"{catering}; alcohol: {yes_no(cfg.alcohol)}; cash handling: {yes_no(cfg.cash_handling)}",
    )
    operations = (
        f"Indoor: {yes_no(cfg.indoor)} | Seated: {yes_no(cfg.seated)} | "
        f"Vehicles: {yes_no(cfg.vehicles)} | Temporary structures: "
        f"{yes_no(cfg.temporary_structures)} | Electrical equipment: "
        f"{yes_no(cfg.electrical_equipment)} | Physical competitions: "
        f"{yes_no(cfg.physical_competitions)}"
    )
    full(11, "Operations", operations)

    next_row = 12
    if cfg.additional_notes:
        full(next_row, "Additional notes", cfg.additional_notes)
        next_row += 1
    for label, value in extra:
        full(next_row, label, value)
        next_row += 1

    tail = doc.add_paragraph()
    set_paragraph_layout(tail, before=2, line=7.5)
    run = tail.add_run("The official Event Management Risk Assessment follows this summary.")
    set_run_font(run, size=7.0, italic=True, colour="555555")


def build_form_document(cfg: EventConfig) -> DocumentType:
    """Build the risk-assessment form document"""
    doc = Document()
    apply_document_defaults(doc)
    set_core_properties(doc, cfg)
    add_form_content(doc, cfg)
    return doc


def build_pack_document(cfg: EventConfig) -> DocumentType:
    """Build a summary page followed by the risk-assessment form"""
    doc = Document()
    apply_document_defaults(doc)
    set_core_properties(doc, cfg)
    add_application_summary_page(doc, cfg)
    doc.add_page_break()
    add_form_content(doc, cfg)
    return doc


def build_document(cfg: EventConfig) -> DocumentType:
    """Build the form-only document for older callers"""
    return build_form_document(cfg)
